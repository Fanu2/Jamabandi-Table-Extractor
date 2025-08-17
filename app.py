import streamlit as st
import pdfplumber
import pandas as pd
from docx import Document
import io

# Add fix_duplicate_columns helper first
def fix_duplicate_columns(df):
    new_cols = []
    counts = {}
    for col in df.columns:
        if col in counts:
            counts[col] += 1
            new_cols.append(f"{col}_{counts[col]}")
        else:
            counts[col] = 0
            new_cols.append(col)
    df.columns = new_cols
    return df


def extract_tables_from_pdf(file):
    tables = []
    with pdfplumber.open(file) as pdf:
        for page in pdf.pages:
            extracted = page.extract_tables()
            for tbl in extracted:
                df = pd.DataFrame(tbl[1:], columns=tbl[0])
                df = fix_duplicate_columns(df)    # keep fix here too
                tables.append(df)
    return tables

def table_to_docx(df, filename="table.docx"):
    doc = Document()
    table = doc.add_table(rows=df.shape[0]+1, cols=df.shape[1])
    for j, column in enumerate(df.columns):
        table.cell(0, j).text = column
    for i in range(df.shape[0]):
        for j in range(df.shape[1]):
            table.cell(i+1, j).text = str(df.iloc[i, j])
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf

st.title("Jamabandi Table Extractor")

uploaded_file = st.file_uploader("Upload a PDF file", type=["pdf"])

if uploaded_file:
    tables = extract_tables_from_pdf(uploaded_file)

    if not tables:
        st.error("No tables found in PDF. Check if PDF has proper grid lines.")
    else:
        st.success(f"Found {len(tables)} table(s) in the PDF.")
        
        # ADD COMBINE MASTER OPTION:
        if len(tables) > 1:
            if st.checkbox("Combine all tables into one master table"):
                master_df = pd.concat(tables, ignore_index=True)
                master_df = fix_duplicate_columns(master_df)
                st.markdown("#### Master Table Preview:")
                st.dataframe(master_df)
                
                # Download master as CSV
                csv_all = master_df.to_csv(index=False).encode('utf-8')
                st.download_button("Download Master CSV", csv_all, file_name="master_table.csv")
                
                # Download master as Excel
                excel_all_buf = io.BytesIO()
                master_df.to_excel(excel_all_buf, index=False)
                excel_all_buf.seek(0)
                st.download_button("Download Master Excel", excel_all_buf.getvalue(), file_name="master_table.xlsx")
                st.markdown("---")

        # Existing table selection block (unchanged)
        index = st.selectbox("Select Table to View or Export",
                             list(range(1, len(tables)+1)),
                             format_func=lambda x: f"Table {x}")

        selected_df = tables[index - 1]
        st.write("Preview of selected table:")
        st.dataframe(selected_df)

        csv_data = selected_df.to_csv(index=False).encode('utf-8')
        st.download_button("Download as CSV", csv_data,
                           file_name=f"table_{index}.csv", mime="text/csv")

        excel_buf = io.BytesIO()
        selected_df.to_excel(excel_buf, index=False)
        excel_buf.seek(0)
        st.download_button("Download as Excel", excel_buf.getvalue(),
                           file_name=f"table_{index}.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        docx_buf = table_to_docx(selected_df, filename=f"table_{index}.docx")
        st.download_button("Download as DOCX", docx_buf,
                           file_name=f"table_{index}.docx",
                           mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
